"""Заполнение исходного шаблона КП средствами python-docx (без зависимости от API ИИ)."""
from pathlib import Path
from datetime import date
from copy import deepcopy
from docx import Document


def rub(value):
    return f"{float(value):,.2f}".replace(",", " ").replace(".", ",")


def _replace_paragraph(paragraph, mapping):
    text = paragraph.text
    for old, new in mapping.items():
        text = text.replace(old, str(new))
    if text != paragraph.text:
        # Preserve paragraph style; placeholders in the supplied template are in standalone runs.
        paragraph.text = text


def _remove_paragraph(paragraph):
    el = paragraph._element
    el.getparent().remove(el)
    paragraph._p = paragraph._element = None


def render_docx(data, out_path, template_path=None, manager="Не назначен", generated_on=None):
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    generated_on = generated_on or date.today()
    if template_path and Path(template_path).exists():
        doc = Document(str(template_path))
    else:
        doc = Document()
        doc.add_paragraph('«Монолит Р»')

    client = data.get("client") or {}
    kp_number = f'КП-{generated_on:%Y%m%d}-{data["letter_id"]}'
    contacts = " · ".join(str(v) for v in (client.get("phone"), client.get("email")) if v) or "Не указаны"
    client_label = " · ".join(str(v) for v in (client.get("name"), client.get("company")) if v) or "Не указан"
    if data.get("delivery_method") == "самовывоз":
        delivery_line = "Самовывоз"
    else:
        d = data.get("delivery") or {}
        if d.get("status") == "рассчитана":
            delivery_line = f'Доставка: {data.get("settlement") or "не указан"}, {d.get("distance_km")} км, зона {d.get("zone")}'
        else:
            delivery_line = f'Доставка: {data.get("settlement") or "не указан"} (стоимость по согласованию)'

    # Fill the supplied template's first metadata table.
    if doc.tables:
        table = doc.tables[0]
        values = {"{{ client }}": client_label, "{{ contacts }}": contacts,
                  "{{ delivery_line }}": delivery_line, "{{ branch }}": data.get("branch") or "Не определён"}
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _replace_paragraph(p, values)

    # Rebuild the item table from its template row, preserving table style and columns.
    if len(doc.tables) >= 2:
        table = doc.tables[1]
        template_row = table.rows[2] if len(table.rows) >= 4 else table.rows[-1]
        row_xml = deepcopy(template_row._tr)
        # Remove Jinja control rows and original sample row.
        for row in list(table.rows[1:]):
            table._tbl.remove(row._tr)
        for idx, item in enumerate(data.get("items", []), 1):
            new_tr = deepcopy(row_xml)
            table._tbl.append(new_tr)
            row = table.rows[-1]
            mark = "акция" if item.get("promo") else ""
            if item.get("on_order_days"):
                mark = f'под заказ {item["on_order_days"]} дн.'
            vals = [str(idx), str(item.get("sku", "")), str(item.get("name", "")) + (f" ({mark})" if mark else ""),
                    str(item.get("unit", "")), str(item.get("qty", "")), rub(item.get("price", 0)), rub(item.get("amount", 0))]
            for cell, value in zip(row.cells, vals):
                cell.text = value
        if not data.get("items"):
            # Keep a clear row when there are no priced items.
            row = table.add_row()
            row.cells[0].merge(row.cells[-1])
            row.cells[0].text = "Нет однозначно определённых позиций для расчёта"

    # Fill totals table.
    if len(doc.tables) >= 3:
        table = doc.tables[2]
        delivery = data.get("delivery") or {}
        if delivery.get("status") == "самовывоз":
            delivery_text = "самовывоз"
        elif delivery.get("status") == "уточнить":
            delivery_text = "по согласованию"
        else:
            delivery_text = f'{delivery.get("vehicle")}, рейсов: {delivery.get("trips")}, зона {delivery.get("zone")}'
        replacements = {
            "{{ goods_total }}": rub(data.get("goods_total", 0)),
            "{{ discount_percent }}": str(data.get("discount_percent", 0)),
            "{{ discount_amount }}": rub(data.get("discount_amount", 0)),
            "{{ delivery_text }}": delivery_text,
            "{{ delivery_cost }}": rub(delivery.get("cost", 0)),
            "{{ total }}": rub(data.get("total", 0)),
            "{{ vat }}": rub(data.get("vat", 0)),
        }
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _replace_paragraph(p, replacements)

    # Remove template's Jinja control paragraphs, then replace remaining placeholders.
    replacements = {"{{ kp_number }}": kp_number, "{{ kp_date }}": generated_on.strftime("%d.%m.%Y"), "{{ manager }}": manager}
    for p in list(doc.paragraphs):
        raw = p.text.strip()
        if raw.startswith("{%p"):
            _remove_paragraph(p)
            continue
        # These headings belong to conditional Jinja sections in the template.
        # Remove them here and rebuild each section exactly once only when non-empty.
        if raw in {"Требует уточнения (в сумму не включено)", "Примечания"}:
            _remove_paragraph(p)
            continue
        # The template's repeated clarify/notes blocks are rebuilt below; remove their marker lines.
        if raw.startswith("• {{ c.request }}") or raw.startswith("• {{ n }}"):
            _remove_paragraph(p)
            continue
        _replace_paragraph(p, replacements)

    # Insert dynamic sections before the template's closing disclaimer, in the
    # same place where the conditional template blocks were intended to appear.
    disclaimer = next((p for p in doc.paragraphs if p.text.startswith("Цены указаны в рублях")), None)
    if data.get("clarify"):
        heading = doc.add_heading("Требует уточнения (в сумму не включено)", level=2)
        if disclaimer is not None:
            disclaimer._p.addprevious(heading._p)
        for item in data["clarify"]:
            paragraph = doc.add_paragraph(
                f'{item.get("request", "Запрос клиента")} — {item.get("reason", "требуется уточнение")}',
                style="List Bullet",
            )
            if disclaimer is not None:
                disclaimer._p.addprevious(paragraph._p)
    unique_notes = list(dict.fromkeys(data.get("notes", [])))
    if unique_notes:
        heading = doc.add_heading("Примечания для менеджера", level=2)
        if disclaimer is not None:
            disclaimer._p.addprevious(heading._p)
        for note in unique_notes:
            paragraph = doc.add_paragraph(str(note), style="List Bullet")
            if disclaimer is not None:
                disclaimer._p.addprevious(paragraph._p)
    doc.save(out_path)
    return out_path
