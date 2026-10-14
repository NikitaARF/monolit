import json
import re
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]

EXPECTED_TOP_KEYS = {
    "letter_id", "client", "delivery_method", "settlement", "branch", "items", "clarify",
    "goods_total", "discount_percent", "discount_base", "discount_amount", "goods_after_discount",
    "weight_kg", "delivery", "total", "vat", "notes",
}
EXPECTED_ITEM_KEYS = {"sku", "name", "unit", "qty", "price", "amount", "promo", "on_order_days"}


def test_all_ten_json_outputs_exist_and_follow_schema():
    results = ROOT / "results"
    for i in range(1, 11):
        path = results / f"kp_{i:02d}.json"
        assert path.exists(), f"Не найден результат: {path.name}"
        data = json.loads(path.read_text(encoding="utf-8"))
        assert set(data) == EXPECTED_TOP_KEYS
        assert data["letter_id"] == f"{i:02d}"
        assert isinstance(data["items"], list)
        assert isinstance(data["clarify"], list)
        assert all(set(item) == EXPECTED_ITEM_KEYS for item in data["items"])
        assert data["discount_percent"] in (0, 3, 5)
        assert data["delivery"]["status"] in ("рассчитана", "самовывоз", "уточнить")


def test_all_ten_docx_outputs_exist_and_have_no_template_tokens():
    results = ROOT / "results"
    for i in range(1, 11):
        path = results / f"kp_{i:02d}.docx"
        assert path.exists(), f"Не найден документ: {path.name}"
        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "{{" not in text
        assert "{%" not in text
        assert re.search(rf"КП-\d{{8}}-{i:02d}", text)
        assert all("{{" not in cell.text and "{%" not in cell.text for table in doc.tables for row in table.rows for cell in row.cells)
